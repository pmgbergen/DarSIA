"""Stage shared protocol workbooks without changing active scientific bindings.

The original readers accept either CSV or ``[workbook, sheet]``. A package keeps
one sheet per run and preserves the source values; activating changed file bytes
is deliberately separate from this storage conversion.
"""
from __future__ import annotations

import csv
from copy import deepcopy
import base64
from dataclasses import dataclass
from datetime import date, datetime, time, timezone
import hashlib
from io import BytesIO, StringIO
import json
import os
from pathlib import Path, PureWindowsPath
import re
import shutil
import tempfile
import tomllib
import uuid

from .storage_paths import machine_local_path, machine_stored_text


KINDS = ("imaging", "injection", "pressure_temperature", "blacklist", "overflow")


@dataclass(frozen=True)
class ProtocolSheet:
    kind: str
    run: str
    columns: tuple[str, ...]
    rows: tuple[tuple[str | None, ...], ...]


@dataclass(frozen=True)
class ProtocolWorkbookPlan:
    root: Path
    runs: tuple[str, ...]
    sheets: tuple[ProtocolSheet, ...]
    # Each run keeps its original folder mapping; the values name a package sheet.
    sections: dict
    source_tables: int
    inputs: tuple[dict, ...] = ()
    configurations: tuple[tuple[str, Path, bytes], ...] = ()


def _text(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _read_table(reference, raw=None):
    """Read tabular values only, keeping decimal and ISO timestamp strings exact."""
    path, sheet = reference if isinstance(reference, (list, tuple)) else (reference, None)
    path = Path(path)
    raw = path.read_bytes() if raw is None else raw
    if path.suffix.casefold() == ".csv":
        if sheet is not None:
            raise ValueError("CSV protocol files cannot name a worksheet")
        values = list(csv.reader(StringIO(raw.decode("utf-8-sig"))))
    elif path.suffix.casefold() == ".xlsx":
        import openpyxl
        workbook = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=True)
        try:
            if sheet is None:
                if len(workbook.sheetnames) != 1:
                    raise ValueError("Select a worksheet for " + str(path))
                sheet = workbook.sheetnames[0]
            values = [list(row) for row in workbook[sheet].iter_rows(values_only=True)]
        finally:
            workbook.close()
    elif path.suffix.casefold() == ".xls":
        import pandas as pd
        frame = pd.read_excel(BytesIO(raw), sheet_name=sheet or 0, dtype=object)
        values = [list(frame.columns)] + frame.where(frame.notna(), None).values.tolist()
    else:
        raise ValueError("Use CSV or Excel protocol tables: " + str(path))
    while values and not any(value is not None and value != "" for value in values[-1]):
        values.pop()
    if not values:
        raise ValueError("The protocol table has no column header: " + str(path))
    width = max((index + 1 for row in values for index, value in enumerate(row)
                 if value is not None and value != ""), default=0)
    columns = tuple(str(value) if value is not None else "" for value in values[0][:width])
    if not columns or any(not name for name in columns) or len(set(columns)) != len(columns):
        raise ValueError("Protocol columns must have distinct non-empty names: " + str(path))
    rows = tuple(tuple(_text(row[index]) if index < len(row) else None
                       for index in range(width)) for row in values[1:]
                 if any(value is not None and value != "" for value in row))
    return columns, rows


def _qualified(value, folder):
    if value is None or not value.strip():
        return value
    path = Path(value)
    if path.is_absolute() or PureWindowsPath(value).is_absolute():
        return value
    normalized = value.replace("\\", "/")
    parts = normalized.split("/")
    base = Path(folder)
    if len(parts) > 1 and parts[0] == base.name:
        base = base.parent
    return machine_stored_text(base / normalized)


def _relative_image_paths(names, rows):
    path_index = next((index for index, name in enumerate(names) if name.casefold() == "path"), None)
    return path_index is not None and any(
        value and not Path(value).is_absolute() and not PureWindowsPath(value).is_absolute()
        for row in rows for value in (row[path_index],))


def plan_protocol_workbooks(series_root, *, config_paths=None, overflow_sources=None):
    """Read declared protocol tables for this series; never scan or read images.

    ``config_paths`` can supply the actual selected configuration chain per run.
    Without it, the same ``config_files`` route used by native series setup runs.
    """
    from .series_paths import config_files
    from .storage_paths import resolve_storage_path
    from darsia.presets.workflows.config.protocols import ProtocolsConfig
    from darsia.presets.workflows.config.utils import _get_section_from_toml
    from darsia.presets.workflows.config.utils import _with_storage_aliases
    from darsia.presets.workflows.config.protocols import _parse_protocol_value

    root = Path(series_root).absolute()
    input_bytes = {}
    def remember(path):
        path = Path(path).absolute()
        if path not in input_bytes:
            input_bytes[path] = path.read_bytes()
        return input_bytes[path]
    manifest = tomllib.loads(remember(root / "series.toml").decode("utf-8-sig"))
    runs = tuple(manifest["runs"]["ids"])
    if (len({run.casefold() for run in runs}) != len(runs)
            or any(not run or len(run) > 31 or re.search(r"[\\/*?:\[\]]", run)
                   or run.startswith("'") or run.endswith("'") for run in runs)):
        raise ValueError("Run IDs must be unique, valid Excel worksheet names")
    cache, sheets, sections, configurations = {}, [], {}, []
    target_runs = {}
    for run in runs:
        paths = list(config_paths[run]) if config_paths is not None else config_files(manifest["series_id"], run)
        for path in paths:
            remember(path)
        path_settings, run_settings = manifest.get("paths", {}), manifest["runs"]
        def configured(name, folder):
            declaration = (run_settings.get(name) or {}).get(run)
            default = "config/runs" if folder == "run_configs" else "config/merged"
            value = declaration or str(Path(path_settings.get(folder, default)) / (run + ".toml"))
            path = Path(value)
            return Path(resolve_storage_path(path if path.is_absolute() else root / path)).absolute(), declaration
        partial, explicit_partial = configured("config_files", "run_configs")
        merged, explicit_merged = configured("merged_files", "merged_configs")
        if not partial.is_file() and config_paths is not None and not explicit_partial:
            partial = next((Path(path).absolute() for path in paths if Path(path).stem == run), Path(paths[-1]).absolute())
        for path, explicit in ((partial, True), (merged, bool(explicit_merged))):
            if not path.is_file():
                if explicit:
                    raise FileNotFoundError(path)
                continue
            if path in target_runs:
                if target_runs[path] != run:
                    raise ValueError("Two runs share the same writable configuration: " + str(path))
                continue
            target_runs[path] = run
            configurations.append((run, path, remember(path)))
        selected = ProtocolsConfig().load(paths)
        raw = _get_section_from_toml(paths, "protocols")
        declarations = {kind: getattr(selected, kind, None) for kind in KINDS}
        if declarations["overflow"] is None and raw.get("overflow"):
            declarations["overflow"] = _parse_protocol_value(raw["overflow"], paths)
        if overflow_sources and run in overflow_sources:
            supplied = _parse_protocol_value(_with_storage_aliases(overflow_sources[run]), paths)
            if declarations["overflow"] is not None and declarations["overflow"] != supplied:
                raise ValueError("The selected run already declares another overflow protocol: " + run)
            declarations["overflow"] = supplied
        sections[run] = {}
        for kind, declaration in declarations.items():
            if declaration is None:
                continue
            mapping = declaration if isinstance(declaration, dict) else {None: declaration}
            columns, merged = [], []
            seen = set()
            for folder, reference in mapping.items():
                key = tuple(map(str, reference)) if isinstance(reference, (list, tuple)) else (str(reference),)
                file = Path(reference[0] if isinstance(reference, (list, tuple)) else reference)
                if kind == "overflow" and file.suffix.casefold() == ".json":
                    key += (run,)
                if key not in cache:
                    raw_table = remember(file)
                    if kind == "overflow" and file.suffix.casefold() == ".json":
                        if isinstance(reference, (list, tuple)):
                            raise ValueError("A reviewed overflow JSON cannot name a worksheet")
                        from .overflow_protocol_import import read_overflow_measurements
                        names, rows = read_overflow_measurements(file, run)
                        cache[key] = tuple(names), tuple(tuple(_text(value) for value in row) for row in rows)
                    else:
                        cache[key] = _read_table(reference, raw_table)
                names, rows = cache[key]
                for name in names:
                    if name not in columns:
                        columns.append(name)
                # Shared full-path sheets are read once. A basename-only source
                # belongs separately to every declared folder using that table.
                group = ((key, str(folder)) if folder is not None and kind in ("imaging", "blacklist")
                         and _relative_image_paths(names, rows) else key)
                if group in seen:
                    continue
                seen.add(group)
                for values in rows:
                    row = dict(zip(names, values))
                    if folder is not None and kind in ("imaging", "blacklist"):
                        path_name = next((name for name in names if name.casefold() == "path"), None)
                        if path_name is not None:
                            row[path_name] = _qualified(row[path_name], folder)
                    merged.append(row)
            sheets.append(ProtocolSheet(kind, run, tuple(columns),
                tuple(tuple(row.get(name) for name in columns) for row in merged)))
            sections[run][kind] = ({str(folder): run for folder in mapping}
                                  if isinstance(declaration, dict) else run)
    inputs = tuple(_descriptor(path, raw) for path, raw in input_bytes.items())
    return ProtocolWorkbookPlan(root, runs, tuple(sheets), sections, len(cache), inputs, tuple(configurations))


def _descriptor(path, raw=None):
    """Bind a file's bytes, read here and recorded in its shared spelling."""
    path = machine_local_path(path).absolute()
    raw = path.read_bytes() if raw is None else raw
    return dict(path=machine_stored_text(path), sha256=hashlib.sha256(raw).hexdigest(), size_bytes=len(raw))


def write_protocol_workbooks(plan, destination):
    """Publish a fresh package; never overwrite sources or activate config files.

    Values are literal text cells. This preserves ISO offsets and decimal source
    precision while the existing readers retain their ordinary numeric parsing.
    """
    import openpyxl
    from openpyxl.cell import WriteOnlyCell

    destination = Path(destination).expanduser().absolute()
    if destination.exists():
        raise FileExistsError("Choose a fresh workbook package folder: " + str(destination))
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = Path(tempfile.mkdtemp(prefix=".protocol-workbooks-", dir=destination.parent))
    paths = {}
    try:
        for kind in KINDS:
            selected = [sheet for sheet in plan.sheets if sheet.kind == kind]
            if not selected:
                continue
            workbook = openpyxl.Workbook(write_only=True)
            for sheet in selected:
                output = workbook.create_sheet(sheet.run)
                output.freeze_panes = "A2"
                for values in (sheet.columns, *sheet.rows):
                    row = []
                    for value in values:
                        cell = WriteOnlyCell(output, value=value)
                        if value is not None:
                            cell.data_type = "s"
                        row.append(cell)
                    output.append(row)
            name = kind + "_protocol.xlsx"
            workbook.save(temporary / name)
            paths[kind] = destination / name
        # Renaming the fresh directory publishes either all workbooks or none.
        temporary.rename(destination)
    except BaseException:
        if not temporary.resolve().is_relative_to(destination.parent.resolve()) or not temporary.name.startswith('.protocol-workbooks-'):
            raise RuntimeError("The private workbook staging folder is outside its output parent")
        shutil.rmtree(temporary)
        raise
    sections = {}
    for run, declarations in plan.sections.items():
        sections[run] = {}
        for kind, selected in declarations.items():
            path = paths[kind]
            locator = (str(path.relative_to(plan.root)) if path.is_relative_to(plan.root)
                       else machine_stored_text(path))
            reference = lambda sheet: [locator, sheet]
            sections[run][kind] = ({folder: reference(sheet) for folder, sheet in selected.items()}
                                   if isinstance(selected, dict) else reference(selected))
    return dict(folder=destination, files=paths, protocols=sections,
                workbooks=tuple(_descriptor(path) for path in paths.values()), activated=False)


def _assert_unchanged(records):
    for expected in records:
        if _descriptor(expected["path"]) != expected:
            raise ValueError("A protocol input changed; preview again: " + expected["path"])


def _exact_replace(path, expected, written):
    """Restore exact captured bytes only if this save still owns the current file."""
    path = Path(path)
    temporary = path.with_name("." + path.name + "." + uuid.uuid4().hex + ".tmp")
    try:
        with temporary.open("xb") as stream:
            stream.write(written)
        if path.read_bytes() != expected:
            raise ValueError("Configuration changed during the workbook update: " + str(path))
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _updated_protocols(path, original, replacement):
    """Change locators only; retain authored folder keys and every other setting."""
    from darsia.presets.workflows.config.utils import _config_path, _with_storage_aliases
    updated = deepcopy(tomllib.loads(original.decode("utf-8-sig")))
    protocols = updated.setdefault("protocols", {})
    for kind, value in replacement.items():
        if kind not in KINDS:
            raise ValueError("Unknown protocol type: " + kind)
        if isinstance(value, dict):
            authored = protocols.get(kind)
            aliases = ({str(_config_path(_with_storage_aliases(key), path)): key for key in authored}
                       if isinstance(authored, dict) else {})
            protocols[kind] = {aliases.get(folder, machine_stored_text(folder)): reference
                               for folder, reference in value.items()}
        else:
            protocols[kind] = deepcopy(value)
    return updated


def activate_protocol_workbooks(plan, package, *, writer=None):
    """Use the normal lossless CAS writer to switch only saved protocol locators.

    One migration record retains exact original configuration bytes. A failed
    multi-file save restores this operation's own writes without replacing a
    concurrent external edit. Protocol data, maps and approvals are untouched.
    """
    from .series_method import _manifest_bytes, _write_manifest
    if package.get("activated"):
        raise ValueError("These protocol workbooks are already active")
    if not plan.configurations or not plan.inputs or not package.get("workbooks"):
        raise ValueError("Preview the saved configurations before activating workbooks")
    _assert_unchanged(plan.inputs)
    _assert_unchanged(package["workbooks"])
    changes = []
    for run, path, original in plan.configurations:
        updated = _updated_protocols(path, original, package["protocols"][run])
        written, parsed = _manifest_bytes(original, updated)
        changes.append((run, path, original, written, parsed))
    record_path = Path(package["folder"]) / "protocol_workbook_migration.json"
    record = dict(format="darsia-protocol-workbook-migration-v1", status="prepared",
        created_utc=datetime.now(timezone.utc).isoformat(), series_root=machine_stored_text(plan.root),
        inputs=list(plan.inputs), workbooks=list(package["workbooks"]),
        configurations=[dict(run_id=run, path=machine_stored_text(path), before={**_descriptor(path, original),
            "bytes_base64": base64.b64encode(original).decode("ascii")},
            after=_descriptor(path, written), protocols_before=tomllib.loads(original.decode("utf-8-sig")).get("protocols", {}),
            protocols_after=parsed["protocols"]) for run, path, original, written, parsed in changes])
    # Exact typed TOML values remain in the byte backup; diagnostic JSON uses
    # ISO strings for TOML's native datetime/date/time values.
    record = json.loads(json.dumps(record, ensure_ascii=False, default=str))
    record_bytes = (json.dumps(record, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    if record_path.exists():
        previous_bytes = record_path.read_bytes()
        previous = json.loads(previous_bytes)
        if (previous.get("status") != "rolled_back"
                or previous.get("configurations") != record["configurations"]
                or previous.get("workbooks") != record["workbooks"]):
            raise FileExistsError("This workbook package already has another migration record")
        _exact_replace(record_path, previous_bytes, record_bytes)
    else:
        with record_path.open("xb") as stream:
            stream.write(record_bytes)
    completed = []
    save = writer or _write_manifest
    try:
        for run, path, original, written, parsed in changes:
            completed.append((path, original, written))
            save(path, original, parsed)
            # The same semantic writer is used for its precomputed bytes above.
            if path.read_bytes() != written:
                raise ValueError("The configuration writer returned unexpected bytes: " + str(path))
        record["status"] = "active"
        final = (json.dumps(record, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
        _exact_replace(record_path, record_bytes, final)
    except BaseException as activation_error:
        failed = []
        for path, original, written in reversed(completed):
            try:
                if path.read_bytes() != original:
                    _exact_replace(path, written, original)
            except Exception as rollback_error:
                failed.append(str(rollback_error))
        record["status"] = "rollback_incomplete" if failed else "rolled_back"
        record["rollback_errors"] = failed
        final = (json.dumps(record, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
        _exact_replace(record_path, record_bytes, final)
        if failed:
            raise RuntimeError("Workbook activation failed; concurrent configuration edits were preserved. "
                               "See the migration record: " + str(record_path)) from activation_error
        raise
    package["activated"] = True
    return dict(folder=package["folder"], migration=record_path, activated=True,
                configurations=[str(path) for _, path, _, _, _ in changes])
