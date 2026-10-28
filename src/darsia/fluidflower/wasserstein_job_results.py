"""Verify real comparison output before whole-run queue publication.

This validates an explicitly supplied owned attempt. It does not discover
results, adopt a scientific result, or replace the queue's publication/terminal
lineage checks. A Jobs snapshot reader must call it only after that lineage is
verified and the original result descriptors resolve on the reading machine.
"""
from __future__ import annotations

from pathlib import Path

from .wasserstein_job_protocol import RESULT_FORMAT, descriptor, validate_wasserstein_claim

OUTPUT_NAMES = {"comparison.json", "wasserstein_distances.csv", "wasserstein_distances.xlsx"}


def _validate_tables(root, row):
    """Require the original single-row exports to represent the JSON scalars."""
    import csv
    import math
    from openpyxl import load_workbook
    from .wasserstein_materialization import RESULT_COLUMNS
    with (root / "wasserstein_distances.csv").open(encoding="utf-8", newline="") as stream:
        records = csv.reader(stream)
        expected = ["" if row[key] is None else str(row[key]) for key in RESULT_COLUMNS]
        if (next(records, None) != list(RESULT_COLUMNS) or next(records, None) != expected
                or next(records, None) is not None):
            raise ValueError("Wasserstein CSV differs from its complete JSON comparison")
    workbook = load_workbook(root / "wasserstein_distances.xlsx", read_only=True, data_only=False)
    try:
        if workbook.sheetnames != ["Wasserstein distances"]:
            raise ValueError("Wasserstein XLSX has a different result sheet")
        records = workbook.worksheets[0].iter_rows()
        header = next(records, ())
        cells = next(records, ())
        if (tuple(cell.value for cell in header) != RESULT_COLUMNS
                or len(cells) != len(RESULT_COLUMNS) or next(records, None) is not None
                or any(cell.data_type == "f" for cell in (*header, *cells))):
            raise ValueError("Wasserstein XLSX has a different comparison table")
        for key, cell in zip(RESULT_COLUMNS, cells):
            actual, expected = cell.value, row[key]
            if expected is None or expected == "":
                matches = actual is None or actual == ""
            elif type(expected) is bool:
                matches = type(actual) is bool and actual == expected
            elif type(expected) in (int, float):
                # Original openpyxl emits numeric cells with 16 significant
                # digits; allow that serialization round trip, never NaN/Inf.
                matches = (type(actual) in (int, float) and math.isfinite(actual)
                           and math.isclose(actual, expected, rel_tol=1e-15, abs_tol=0.0))
            else:
                matches = type(actual) is str and actual == expected
            if not matches:
                raise ValueError("Wasserstein XLSX differs from its JSON scalar: " + key)
    finally:
        workbook.close()


def validate_wasserstein_result(result, *, task, job, output_root):
    """Hash-check every declared product and its actual scalar comparison row."""
    from .pipeline.job_queue import descriptor as file_descriptor
    from .wasserstein_jobs import read_bound
    from .wasserstein_job_runner import validate_solver_row
    validate_wasserstein_claim(task, job)
    if (not isinstance(result, dict) or result.get("format") != RESULT_FORMAT
            or result.get("complete") is not True or result.get("pipeline") != task["pipeline"]
            or result.get("run") != task["run"] or result.get("backend") != "cpu"
            or result.get("precision") != "float64" or result.get("comparisons") != 1
            or result.get("frames") != 0 or result.get("comparison") != task["comparison"]
            or result.get("wasserstein_contract_sha256") != task["wasserstein_contract_sha256"]
            or result.get("wasserstein_plan") != task["wasserstein_plan"]
            or result.get("worker_references_closed") is not True
            or result.get("validation", {}).get("state") != "not_scored"):
        raise ValueError("Wasserstein result does not identify the actual completed comparison")
    files = result.get("files")
    if not isinstance(files, dict) or set(files) != OUTPUT_NAMES:
        raise ValueError("A completed Wasserstein comparison needs its JSON, CSV and XLSX outputs")
    root = Path(output_root).resolve()
    for filename in sorted(OUTPUT_NAMES):
        source = descriptor(files[filename])
        path = Path(source["path"])
        if path.resolve() != root / filename or path.name != filename:
            raise ValueError("Wasserstein output is outside the explicitly owned attempt")
        if file_descriptor(path) != source:
            raise ValueError("A Wasserstein output changed after computation: " + filename)
    row = read_bound(files["comparison.json"])
    validate_solver_row(row, task["comparison"])
    if (result.get("solver_status") != row["status"]
            or result.get("solver_converged") != row["solver_converged"]):
        raise ValueError("Wasserstein result summary differs from its actual scalar output")
    _validate_tables(root, row)
    for filename in sorted(OUTPUT_NAMES):
        if file_descriptor(root / filename) != descriptor(files[filename]):
            raise ValueError("A Wasserstein output changed during result validation: " + filename)
    return result
