"""Shared protocol storage preserves the original readers' scientific inputs."""
from pathlib import Path
import base64
from datetime import datetime
import csv
from io import StringIO
from types import SimpleNamespace

import openpyxl
import pytest
import toml
import json

from darsia import ImagingProtocol, InjectionProtocol, PressureTemperatureProtocol
from darsia.fluidflower.protocol_workbooks import (
    activate_protocol_workbooks, plan_protocol_workbooks, write_protocol_workbooks,
)
from darsia.presets.workflows.config.protocols import ProtocolsConfig


def put_table(path, header, rows):
    data = StringIO(newline="")
    writer = csv.writer(data)
    writer.writerow(header)
    writer.writerows(rows)
    path.write_text(data.getvalue(), encoding="utf-8")
    return str(path)


@pytest.fixture
def source(tmp_path):
    root = tmp_path / "series"
    root.mkdir()
    (root / "series.toml").write_text(toml.dumps(dict(schema_version=1,
        series_id="test", runs=dict(ids=["r1"]))), encoding="utf-8")
    first, second = root / "photos1", root / "photos2"
    first.mkdir(); second.mkdir()
    imaging = {}
    for index, folder in enumerate((first, second)):
        imaging[str(folder)] = put_table(root / f"imaging{index}.csv",
            ["path", "datetime", "source_sha256", "operator_note"],
            [["same.JPG", f"2023-01-01T0{index}:00:00.123456+01:00", "0000123", "=literal"]])
    injection = put_table(root / "injection.csv",
        ["id", "location_x", "location_y", "start", "end", "rate_sccm", "density kg/m3"],
        [[1, "0.5", "0.02", "2023-01-01T01:00:00+01:00", "2023-01-01T02:00:00+01:00",
          "0.123456789012345678", "1.872"]])
    pressure = put_table(root / "pressure.csv",
        ["datetime", "pressure_bar", "temperature_celsius", "pressure_gradient_bar", "temperature_gradient_celsius"],
        [["2023-01-01T00:00:00+01:00", "0.998798", "23.0", "0.0", "0.0"],
         ["2023-01-02T00:00:00+01:00", "0.998798", "23.0", "0.0", "0.0"]])
    blacklist = put_table(root / "blacklist.csv", ["path"], [["same.JPG"]])
    protocols = dict(imaging=imaging, injection=injection, pressure_temperature=pressure,
        blacklist={str(first): blacklist}, imaging_mode="Load from CSV")
    config = root / "r1.toml"
    config.write_text(toml.dumps(dict(data=dict(folders=[str(first), str(second)]), protocols=protocols)), encoding="utf-8")
    return SimpleNamespace(root=root, paths={"r1": [config]}, config=config,
        first=first, second=second, protocols=protocols)


def test_shared_run_sheet_preserves_reader_outputs_and_folder_blacklist(source):
    plan = plan_protocol_workbooks(source.root, config_paths=source.paths)
    before = {path: path.read_bytes() for path in source.root.iterdir() if path.is_file()}
    package = write_protocol_workbooks(plan, source.root / "protocols" / "shared")
    assert package["activated"] is False
    assert len(package["files"]) == 4
    assert "overflow" not in package["files"]
    assert all(path.read_bytes() == data for path, data in before.items())
    declarations = package["protocols"]["r1"]
    reference = lambda value: (source.root / value[0], value[1])
    for folder, original in source.protocols["imaging"].items():
        new = ImagingProtocol(reference(declarations["imaging"][folder]))
        old = ImagingProtocol(Path(original))
        assert new.get_datetime(Path(folder) / "same.JPG") == old.get_datetime(Path(folder) / "same.JPG")
    new_blacklist = ImagingProtocol(reference(declarations["imaging"][str(source.first)]),
        reference(declarations["blacklist"][str(source.first)]))
    assert new_blacklist.is_blacklisted(source.first / "same.JPG")
    assert not new_blacklist.is_blacklisted(source.second / "same.JPG")
    old_injection = InjectionProtocol(Path(source.protocols["injection"]))
    new_injection = InjectionProtocol(reference(declarations["injection"]))
    assert new_injection.df.equals(old_injection.df)
    old_pt = PressureTemperatureProtocol(Path(source.protocols["pressure_temperature"]))
    new_pt = PressureTemperatureProtocol(reference(declarations["pressure_temperature"]))
    assert new_pt.df.equals(old_pt.df)


def test_literal_cells_and_all_source_columns_survive(source):
    plan = plan_protocol_workbooks(source.root, config_paths=source.paths)
    package = write_protocol_workbooks(plan, source.root / "package")
    workbook = openpyxl.load_workbook(package["files"]["imaging"], read_only=True, data_only=False)
    sheet = workbook["r1"]
    rows = list(sheet.values)
    assert rows[0] == ("path", "datetime", "source_sha256", "operator_note")
    assert rows[1][1:] == ("2023-01-01T00:00:00.123456+01:00", "0000123", "=literal")
    assert sheet["D2"].data_type == "s"
    workbook.close()
    workbook = openpyxl.load_workbook(package["files"]["injection"], read_only=True)
    assert workbook["r1"]["F2"].value == "0.123456789012345678"
    workbook.close()


def test_existing_package_and_sources_never_overwritten(source):
    plan = plan_protocol_workbooks(source.root, config_paths=source.paths)
    existing = source.root / "existing"
    existing.mkdir()
    sentinel = existing / "imaging_protocol.xlsx"
    sentinel.write_bytes(b"keep existing")
    with pytest.raises(FileExistsError):
        write_protocol_workbooks(plan, existing)
    assert sentinel.read_bytes() == b"keep existing"


def test_optional_overflow_is_copied_without_unit_or_time_inference(source):
    raw = toml.loads(source.config.read_text(encoding="utf-8"))
    raw["protocols"]["overflow"] = put_table(source.root / "overflow.csv",
        ["datetime", "free_gas_volume_ml", "reviewed"],
        [["2023-01-01T01:00:00+01:00", "1.23456789012345678", "true"]])
    source.config.write_text(toml.dumps(raw), encoding="utf-8")
    plan = plan_protocol_workbooks(source.root, config_paths=source.paths)
    package = write_protocol_workbooks(plan, source.root / "package")
    workbook = openpyxl.load_workbook(package["files"]["overflow"], read_only=True)
    assert list(workbook["r1"].values) == [
        ("datetime", "free_gas_volume_ml", "reviewed"),
        ("2023-01-01T01:00:00+01:00", "1.23456789012345678", "true")]
    workbook.close()


def test_two_runs_share_one_workbook_without_duplicate_common_rows(source):
    manifest = toml.loads((source.root / "series.toml").read_text())
    manifest["runs"]["ids"].append("r2")
    (source.root / "series.toml").write_text(toml.dumps(manifest), encoding="utf-8")
    second_config = source.root / "r2.toml"
    second_config.write_bytes(source.config.read_bytes())
    paths = dict(source.paths, r2=[second_config])
    plan = plan_protocol_workbooks(source.root, config_paths=paths)
    package = write_protocol_workbooks(plan, source.root / "package")
    workbook = openpyxl.load_workbook(package["files"]["imaging"], read_only=True)
    assert workbook.sheetnames == ["r1", "r2"]
    assert len(list(workbook["r1"].values)) == 3
    workbook.close()


def test_shared_basename_table_retains_both_declared_folders(source):
    raw = toml.loads(source.config.read_text(encoding="utf-8"))
    raw["protocols"]["imaging"][str(source.second)] = raw["protocols"]["imaging"][str(source.first)]
    source.config.write_text(toml.dumps(raw), encoding="utf-8")
    plan = plan_protocol_workbooks(source.root, config_paths=source.paths)
    package = write_protocol_workbooks(plan, source.root / "package")
    reference = package["protocols"]["r1"]["imaging"][str(source.first)]
    current = ImagingProtocol((source.root / reference[0], reference[1]))
    assert current.get_datetime(source.first / "same.JPG") == current.get_datetime(source.second / "same.JPG")


def test_activate_changes_only_locators_in_run_and_merged_configs(source):
    authored = toml.loads(source.config.read_text())
    authored["protocols"]["start_reference_fixed"] = datetime(2023, 1, 1, 0, 0)
    source.config.write_text(toml.dumps(authored), encoding="utf-8")
    merged = source.root / "config/merged/r1.toml"
    merged.parent.mkdir(parents=True)
    merged.write_bytes(b"# existing merged configuration\r\n" + source.config.read_bytes())
    originals = {path: path.read_bytes() for path in (source.config, merged)}
    old = ProtocolsConfig().load(source.config)
    plan = plan_protocol_workbooks(source.root, config_paths=source.paths)
    package = write_protocol_workbooks(plan, source.root / "package")
    result = activate_protocol_workbooks(plan, package)
    assert result["activated"] and package["activated"]
    assert set(result["configurations"]) == {str(path) for path in originals}
    for path, original in originals.items():
        first = toml.loads(original.decode())
        after = toml.loads(path.read_text())
        for kind in package["protocols"]["r1"]:
            first["protocols"].pop(kind, None)
            after["protocols"].pop(kind, None)
        assert first == after
        current = ProtocolsConfig().load(path)
        assert InjectionProtocol(current.injection).df.equals(InjectionProtocol(old.injection).df)
        assert PressureTemperatureProtocol(current.pressure_temperature).df.equals(PressureTemperatureProtocol(old.pressure_temperature).df)
        for folder in current.imaging:
            assert ImagingProtocol(current.imaging[folder]).get_datetime(folder / "same.JPG") == ImagingProtocol(old.imaging[folder]).get_datetime(folder / "same.JPG")
        blocked = ImagingProtocol(current.imaging[source.first], current.blacklist[source.first])
        assert blocked.is_blacklisted(source.first / "same.JPG")
        assert not blocked.is_blacklisted(source.second / "same.JPG")
    record = json.loads(result["migration"].read_text())
    assert record["status"] == "active"
    assert len(record["configurations"]) == 2
    for entry in record["configurations"]:
        assert base64.b64decode(entry["before"]["bytes_base64"]) == originals[Path(entry["path"])]


@pytest.mark.parametrize("changed", ["source", "configuration", "workbook"])
def test_activation_refuses_changed_inputs_before_writing_any_config(source, changed):
    original = source.config.read_bytes()
    plan = plan_protocol_workbooks(source.root, config_paths=source.paths)
    package = write_protocol_workbooks(plan, source.root / "package")
    target = {"source": Path(source.protocols["injection"]), "configuration": source.config,
              "workbook": package["files"]["imaging"]}[changed]
    target.write_bytes(target.read_bytes() + b"\n")
    with pytest.raises(ValueError, match="changed"):
        activate_protocol_workbooks(plan, package)
    assert source.config.read_bytes() == original + (b"\n" if changed == "configuration" else b"")
    assert not (package["folder"] / "protocol_workbook_migration.json").exists()


def test_failed_activation_rolls_back_exact_bytes_and_can_retry(source):
    from darsia.fluidflower.series_method import _write_manifest
    merged = source.root / "config/merged/r1.toml"
    merged.parent.mkdir(parents=True)
    merged.write_bytes(b"# preserve exact original whitespace\r\n" + source.config.read_bytes())
    original = {path: path.read_bytes() for path in (source.config, merged)}
    plan = plan_protocol_workbooks(source.root, config_paths=source.paths)
    package = write_protocol_workbooks(plan, source.root / "package")
    def fail_second(path, before, updated):
        if path == merged:
            raise OSError("simulated write failure")
        return _write_manifest(path, before, updated)
    with pytest.raises(OSError, match="simulated"):
        activate_protocol_workbooks(plan, package, writer=fail_second)
    assert all(path.read_bytes() == raw for path, raw in original.items())
    record = json.loads((package["folder"] / "protocol_workbook_migration.json").read_text())
    assert record["status"] == "rolled_back"
    assert not package["activated"]
    assert activate_protocol_workbooks(plan, package)["activated"]


def test_rollback_preserves_concurrent_external_edits(source):
    from darsia.fluidflower.series_method import _write_manifest
    merged = source.root / "config/merged/r1.toml"
    merged.parent.mkdir(parents=True)
    merged.write_bytes(source.config.read_bytes())
    plan = plan_protocol_workbooks(source.root, config_paths=source.paths)
    package = write_protocol_workbooks(plan, source.root / "package")
    def concurrent_second(path, before, updated):
        if path == merged:
            source.config.write_bytes(b"# concurrent external edit\n" + source.config.read_bytes())
            raise OSError("second save failed")
        return _write_manifest(path, before, updated)
    with pytest.raises(RuntimeError, match="concurrent configuration edits were preserved"):
        activate_protocol_workbooks(plan, package, writer=concurrent_second)
    assert source.config.read_bytes().startswith(b"# concurrent external edit\n")
    record = json.loads((package["folder"] / "protocol_workbook_migration.json").read_text())
    assert record["status"] == "rollback_incomplete"
    assert len(record["rollback_errors"]) == 1


def test_explicit_overflow_source_does_not_replace_an_existing_declaration(source):
    first = put_table(source.root / "overflow1.csv", ["datetime", "weight_g"], [["2023-01-01T01:00:00+01:00", 1]])
    second = put_table(source.root / "overflow2.csv", ["datetime", "weight_g"], [["2023-01-01T01:00:00+01:00", 2]])
    plan = plan_protocol_workbooks(source.root, config_paths=source.paths, overflow_sources={"r1": first})
    assert any(sheet.kind == "overflow" for sheet in plan.sheets)
    raw = toml.loads(source.config.read_text())
    raw["protocols"]["overflow"] = first
    source.config.write_text(toml.dumps(raw), encoding="utf-8")
    with pytest.raises(ValueError, match="already declares another"):
        plan_protocol_workbooks(source.root, config_paths=source.paths, overflow_sources={"r1": second})
